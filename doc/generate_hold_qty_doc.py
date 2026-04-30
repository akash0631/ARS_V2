"""
Generates HOLD_QTY_Lifecycle.docx — layman-language walkthrough of the
end-to-end HOLD_QTY process from inward GRN to store delivery, including:

  - the 8-stage narrative with one worked numerical example
  - a visual flowchart (boxes + arrows in native Word shapes/tables)
  - a state-of-each-table cheat-sheet
  - the four common confusions answered
  - the gap-fix recommendations (Option A / Option B) with code blocks

Style mirrors doc/generate_ars_guide.py.
Run:  python doc/generate_hold_qty_doc.py
Output: doc/HOLD_QTY_Lifecycle.docx
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ─── Style helpers (cribbed from generate_ars_guide.py) ─────────────────────
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY_TEXT = RGBColor(0x55, 0x55, 0x55)


def shade_cell(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color="BFBFBF", sz="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), sz)
        b.set(qn("w:color"), color)
        tcBorders.append(b)
    tc_pr.append(tcBorders)


def add_hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "808080")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_para(doc, text, *, bold=False, italic=False, size=11, color=None,
             align=None, space_after=4, mono=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if mono:
        run.font.name = "Consolas"
    if color:
        run.font.color.rgb = RGBColor(*color)
    return p


def add_h(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY
    return h


def add_bullets(doc, items, *, style="List Bullet", size=11):
    for it in items:
        p = doc.add_paragraph(style=style)
        run = p.add_run(it)
        run.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(2)


def add_numbered(doc, items, size=11):
    for it in items:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(it)
        run.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(2)


def add_callout(doc, label, body, fill_hex):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = True
    cell = tbl.cell(0, 0)
    shade_cell(cell, fill_hex)
    set_cell_borders(cell)
    p1 = cell.paragraphs[0]
    r1 = p1.add_run(label + "  ")
    r1.bold = True
    r1.font.size = Pt(11)
    r1.font.color.rgb = NAVY
    r2 = p1.add_run(body)
    r2.font.size = Pt(11)
    doc.add_paragraph()


def info_box(doc, body):
    add_callout(doc, "IN PLAIN ENGLISH →", body, "DDEBF7")


def warn_box(doc, body):
    add_callout(doc, "WATCH OUT →", body, "FCE4D6")


def tip_box(doc, body):
    add_callout(doc, "TIP →", body, "E2EFDA")


def code_block(doc, text, lang_label=""):
    """Monospace shaded block for SQL / Python snippets."""
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    shade_cell(cell, "F2F2F2")
    set_cell_borders(cell, color="DCDCDC")
    if lang_label:
        p = cell.paragraphs[0]
        r = p.add_run(lang_label)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = GREY_TEXT
        cell.add_paragraph()
    code_p = cell.add_paragraph()
    code_p.paragraph_format.space_after = Pt(0)
    for line in text.splitlines():
        r = code_p.add_run(line + "\n")
        r.font.name = "Consolas"
        r.font.size = Pt(9)
    doc.add_paragraph()


def make_table(doc, headers, rows, col_widths_in=None,
               header_fill="1F3A5F", header_color=(255, 255, 255), zebra=True):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.autofit = True
    if col_widths_in:
        for i, w in enumerate(col_widths_in):
            for c in tbl.columns[i].cells:
                c.width = Inches(w)
    # Header row
    for i, h in enumerate(headers):
        c = tbl.cell(0, i)
        shade_cell(c, header_fill)
        set_cell_borders(c)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(*header_color)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    # Body rows
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = tbl.cell(ri + 1, ci)
            if zebra and (ri % 2 == 1):
                shade_cell(c, "F7F7F7")
            set_cell_borders(c)
            p = c.paragraphs[0]
            r = p.add_run(str(val))
            r.font.size = Pt(10)
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    doc.add_paragraph()


# ─── Flowchart helpers ──────────────────────────────────────────────────────
def flow_node(doc, title, subtitle=None, fill="DDEBF7", border="4F81BD"):
    """A single-cell shaded box used as a flowchart node."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(4.5)
    shade_cell(cell, fill)
    set_cell_borders(cell, color=border, sz="8")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = NAVY
    if subtitle:
        sp = cell.add_paragraph()
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sp.paragraph_format.space_after = Pt(0)
        sr = sp.add_run(subtitle)
        sr.italic = True
        sr.font.size = Pt(9)
        sr.font.color.rgb = GREY_TEXT


def flow_arrow(doc, label="▼"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label)
    r.font.size = Pt(14)
    r.bold = True
    r.font.color.rgb = NAVY


def flow_branch(doc, left_title, left_subtitle, right_title, right_subtitle,
                left_fill="E2EFDA", right_fill="FFF2CC",
                left_border="70AD47", right_border="BF8F00"):
    """A two-column branch in the flowchart."""
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for ci, (title, subtitle, fill, border) in enumerate([
        (left_title, left_subtitle, left_fill, left_border),
        (right_title, right_subtitle, right_fill, right_border),
    ]):
        cell = tbl.cell(0, ci)
        cell.width = Inches(2.7)
        shade_cell(cell, fill)
        set_cell_borders(cell, color=border, sz="8")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = NAVY
        sp = cell.add_paragraph()
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sp.paragraph_format.space_after = Pt(0)
        sr = sp.add_run(subtitle)
        sr.italic = True
        sr.font.size = Pt(9)
        sr.font.color.rgb = GREY_TEXT


# ─── Document body ──────────────────────────────────────────────────────────
def build_doc(out_path: Path):
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    # Default font
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # ── Title block ────────────────────────────────────────────────────────
    title = doc.add_heading("HOLD_QTY — End-to-End Lifecycle", level=0)
    for r in title.runs:
        r.font.color.rgb = NAVY

    add_para(doc,
             "Inward GRN → Store delivery; the complete walkthrough",
             italic=True, size=12, color=(0x55, 0x55, 0x55),
             align=WD_ALIGN_PARAGRAPH.LEFT)
    add_para(doc,
             "ARS V2 Retail Auto Replenishment System  |  generated 2026-04-29",
             italic=True, size=9, color=(0x80, 0x80, 0x80))
    add_hr(doc)

    # ── TL;DR ──────────────────────────────────────────────────────────────
    add_h(doc, "TL;DR (one-page cheat sheet)", level=1)
    add_bullets(doc, [
        "HOLD_QTY only exists for TBL options (brand-new listings). RL and TBC always have HOLD_QTY = 0.",
        "ALLOC_QTY (= SHIP_QTY) ships to the store via BDC. HOLD_QTY stays at the RDC as a lead-time reserve.",
        "OPT_MBQ_WH = OPT_MBQ + (rate × hold_days). The difference becomes HOLD_QTY when the pool covers the full demand.",
        "Per allocation row: pool_taken is split into ship-first, hold-residual. Ship demand wins on shortfall.",
        "ARS_NL_TBL_HOLD_TRACKING records HOLD_QTY at the (WERKS, VAR_ART, SZ) grain.",
        "Tracker decrement (unconditional): HOLD_REM = MAX(0, HOLD_REM − ALLOC_QTY); row closes when HOLD_REM hits 0.",
        "Physical source when ALLOC_QTY > HOLD_REM: take HOLD_REM from hold first, then if FNL_Q_REM covers the difference take it from there; otherwise the gap is unfulfilled.",
        "✓ Gap closed: MSA Step 6.5 now subtracts open HOLD_REM from STK_QTY before computing FNL_Q (alongside PEND_QTY).",
    ])

    # ── Section 1 ──────────────────────────────────────────────────────────
    add_h(doc, "1. What is HOLD_QTY?", level=1)
    info_box(doc,
             "When a brand-new article first lists at a store (OPT_TYPE = TBL = "
             "'to-be-listed'), the planner doesn't just send enough stock for one "
             "selling cycle. It sends one cycle's worth (OPT_MBQ) PLUS a reserve at "
             "the warehouse to cover the lead time before the next batch arrives. "
             "The 'send to store now' portion is ALLOC_QTY (= SHIP_QTY). The 'keep "
             "at the RDC for next time' portion is HOLD_QTY. HOLD only ever exists "
             "for TBL options. RL and TBC always have HOLD = 0.")
    add_para(doc,
             "Think of it like restocking a vending machine: ALLOC_QTY is what you "
             "load into the machine today; HOLD_QTY is the box of extras you leave "
             "in the cupboard so you don't have to drive back to the warehouse "
             "before next week's refill.", size=11)

    # ── Section 2: the worked example ──────────────────────────────────────
    doc.add_page_break()
    add_h(doc, "2. The Whole Story — Inward → Outward", level=1)
    add_para(doc,
             "We follow 200 units of a brand-new article (GEN_ART=1116111940, "
             "colour LT_PST, size M, VAR_ART=8801) from the moment they land at "
             "warehouse DH24 until they reach the store shelf. Five stores will "
             "need this article on day 1; a sixth shows up on day 3.")

    add_h(doc, "Stage 1 — Stock arrives at the RDC (Inward)", level=2)
    add_bullets(doc, [
        "Day 0 evening: warehouse DH24 receives a 200-unit GRN from SAP.",
        "ET_STORE_STOCK now shows: RDC=DH24, VAR_ART=8801, SZ=M, V01_FRESH = 200.",
        "Nothing has been promised to anyone yet. MASTER_ALC_PEND has 0 rows for this article.",
    ])

    add_h(doc, "Stage 2 — MSA computes the shippable pool", level=2)
    add_para(doc,
             "Day 1 morning, planner clicks MSA Calculate. The 9-step algorithm "
             "produces the FNL_Q value used by the allocator.")
    make_table(doc,
               ["Step", "What it does", "Result for our 200 units"],
               [
                   ["5", "Pivot SLOCs, sum stock", "STK_QTY = 200"],
                   ["6", "Subtract MASTER_ALC_PEND", "PEND_QTY = 0 (nothing pending)"],
                   ["7", "FNL_Q = max(STK − PEND, 0)", "FNL_Q = 200"],
               ],
               col_widths_in=[0.6, 3.0, 2.4])
    add_para(doc,
             "ARS_MSA_VAR_ART row: RDC=DH24, VAR_ART=8801, SZ=M, FNL_Q=200. "
             "This 200 is the pool the allocator may draw from.")

    add_h(doc, "Stage 3 — Listing classifies and computes demand", level=2)
    add_para(doc,
             "For each of the 5 stores (HN05–HN09), Part 3.6 sees STK_TTL = 0 (no "
             "store stock yet) and tags OPT_TYPE = 'TBL'. Part 4c then computes "
             "demand. Suppose ACS_D = 2 units/day, ALC_D = 7 days, hold_days = 3, "
             "sale rate = 2/day:")
    code_block(doc,
               "OPT_MBQ    = ACS_D + rate × ALC_D            = 2 + 2 × 7         = 16  (cycle stock)\n"
               "OPT_MBQ_WH = ACS_D + rate × (ALC_D+hold_days) = 2 + 2 × (7+3)    = 22  (cycle + buffer)\n"
               "SZ_SHIP_REQ ≈ OPT_MBQ - STK_TTL              = 16 - 0           = 16  per store\n"
               "SZ_POOL_REQ ≈ OPT_MBQ_WH - STK_TTL           = 22 - 0           = 22  per store\n"
               "HOLD buffer  = SZ_POOL_REQ - SZ_SHIP_REQ                         =  6  per store",
               lang_label="Demand math")
    add_para(doc,
             "Five stores × 22 = 110 units of total pool demand. 80 will ship, "
             "30 will be held.")

    add_h(doc, "Stage 4 — Allocator splits SHIP and HOLD", level=2)
    add_para(doc,
             "Pool starts at FNL_Q_REM = 200. Allocator walks the 5 stores in "
             "store-rank order. Per-row split logic from rule_engine.py:1281-1303:")
    code_block(doc,
               "pool_taken = MIN(SZ_POOL_REQ, FNL_Q_REM - prev_pool_demand)\n\n"
               "if OPT_TYPE = 'TBL' and pool_taken > SZ_SHIP_REQ:\n"
               "    SHIP_QTY = SZ_SHIP_REQ          -- capped at must-ship demand\n"
               "    HOLD_QTY = pool_taken - SZ_SHIP_REQ\n"
               "else:\n"
               "    SHIP_QTY = pool_taken           -- ship everything (RL/TBC or pool short)\n"
               "    HOLD_QTY = 0",
               lang_label="rule_engine.py allocation split")
    make_table(doc,
               ["Store", "pool_taken", "SHIP_QTY (= ALLOC_QTY)", "HOLD_QTY", "FNL_Q_REM after"],
               [
                   ["HN05 (rank 1)", "22", "16", "6", "178"],
                   ["HN06 (rank 2)", "22", "16", "6", "156"],
                   ["HN07 (rank 3)", "22", "16", "6", "134"],
                   ["HN08 (rank 4)", "22", "16", "6", "112"],
                   ["HN09 (rank 5)", "22", "16", "6", "90"],
               ],
               col_widths_in=[1.6, 1.0, 1.7, 1.0, 1.4])
    add_para(doc,
             "Total: 80 ship, 30 hold, 90 left in the pool unallocated.",
             bold=True)
    add_para(doc,
             "Part 8.5 then sets OPT_STATUS = 'NL' for each row (since "
             "STK + ALLOC ≥ threshold × ACS_D, the option is now 'new listing, "
             "adequate after alloc').")

    add_h(doc, "Stage 5 — HOLD is recorded in the tracker", level=2)
    add_para(doc,
             "Part 8.6 fires. Step B inserts 5 fresh rows in "
             "ARS_NL_TBL_HOLD_TRACKING (one per store) because each row has "
             "HOLD_QTY > 0 and OPT_STATUS in ('NL','TBL'):")
    code_block(doc,
               "WERKS=HN05, VAR_ART=8801, SZ=M, OPT_STATUS=NL,\n"
               "  HOLD_QTY_INITIAL=6, HOLD_REM=6, IS_CLOSED=0, LISTED_DATE=now\n"
               "WERKS=HN06, ...                                    HOLD_REM=6\n"
               "WERKS=HN07, ...                                    HOLD_REM=6\n"
               "WERKS=HN08, ...                                    HOLD_REM=6\n"
               "WERKS=HN09, ...                                    HOLD_REM=6",
               lang_label="ARS_NL_TBL_HOLD_TRACKING after Stage 5")
    add_para(doc, "30 units are now earmarked for next cycle of these 5 stores.")

    add_h(doc, "Stage 6 — Outbound dispatch (Outward, ship side)", level=2)
    add_bullets(doc, [
        "Day 1 evening: planner exports BDC. ALLOC_QTY = 80 units total (16 × 5) goes to the dispatch file. The 30 hold units are NOT in the dispatch — they stay at DH24.",
        "Day 2 morning: warehouse picks the 80, SAP issues outbound deliveries, MASTER_ALC_PEND gets populated with 80 units pending.",
        "Day 3: 80 units reach the stores. ET_STORE_STOCK refresh shows each store with 16 units; MASTER_ALC_PEND clears as deliveries are confirmed.",
    ])
    tip_box(doc,
            "After Stage 6 the 200 inward units are split as: 80 at stores, "
            "30 physically at DH24 earmarked as HOLD for HN05–HN09, "
            "and 90 physically at DH24 unallocated free-pool.")

    add_h(doc, "Stage 7 — Next cycle: HOLD is consumed", level=2)
    add_para(doc,
             "Day 7, the planner runs MSA + Listing again. Each of the 5 stores "
             "has sold ~4 units, so STK_TTL = 12 per store. The article is now "
             "stocked, so its OPT_TYPE flips to RL (or TBC). HOLD_QTY for this "
             "round is 0 (HOLD only fires for TBL).")
    add_para(doc,
             "Allocator computes SZ_SHIP_REQ = 16 − 12 = 4 per store. Each store "
             "is allocated ALLOC_QTY = 4. ARS_ALLOC_WORKING rows: ALLOC_QTY=4, "
             "HOLD_QTY=0 each. Part 8.6 then fires (cycle 2: ALLOC=4 ≤ HOLD_REM=6, "
             "so the easy case):")
    add_bullets(doc, [
        "consumed_from_hold = MIN(ALLOC_QTY, HOLD_REM) = MIN(4, 6) = 4. The store ships entirely from its held reserve.",
        "extra_needed = ALLOC_QTY − consumed_from_hold = 0. No need to touch FNL_Q_REM.",
        "Tracker decrement: HOLD_REM = MAX(0, HOLD_REM − ALLOC_QTY) = MAX(0, 6 − 4) = 2. Row stays open.",
        "STEP A.5 doesn't match (no new HOLD_QTY > 0 in this run).",
        "STEP B doesn't match (tracker row exists already).",
    ])
    add_para(doc,
             "Tracker after cycle 2: each row has HOLD_REM=2, IS_CLOSED=0. "
             "20 hold units have been consumed (shipped to stores), 10 still earmarked.")

    add_para(doc,
             "Day 14, cycle 3. Each store needs 4 again — the interesting case "
             "where ALLOC_QTY > HOLD_REM. Now we have to ask where the extra "
             "units come from:", bold=False)
    add_numbered(doc, [
        "Take from hold first: consumed_from_hold = MIN(ALLOC_QTY, HOLD_REM) = MIN(4, 2) = 2.",
        "Difference still owed: extra_needed = ALLOC_QTY − consumed_from_hold = 4 − 2 = 2.",
        "Because ALLOC_QTY > HOLD_REM, check FNL_Q_REM (the free pool the allocator saw at allocation time):",
    ])
    add_bullets(doc, [
        "If FNL_Q_REM ≥ extra_needed → the difference (2 units) comes from the free pool. The store receives the full ALLOC_QTY = 4: 2 from hold + 2 from pool.",
        "If FNL_Q_REM is empty or short → only MIN(ALLOC_QTY, HOLD_REM) = 2 units actually ship. The shortfall is unfulfilled this cycle.",
    ])
    add_para(doc,
             "In our example FNL_Q_REM has plenty of stock (the original pool of "
             "200 has only released 80 + 20 so far), so each store gets 4 delivered "
             "as 2-from-hold + 2-from-pool.")
    add_para(doc,
             "Tracker decrement is unconditional and runs the same SQL regardless: "
             "HOLD_REM = MAX(0, HOLD_REM − ALLOC_QTY) = MAX(0, 2 − 4) = 0 → "
             "IS_CLOSED = 1, CLOSED_DATE = now. All 5 tracker rows close. The hold "
             "pool is fully consumed.")
    info_box(doc,
             "Why the conditional logic matters. The SQL in Part 8.6 Step A is "
             "simple — it just clamps HOLD_REM at 0. But the PHYSICAL accounting "
             "depends on the pool: when ALLOC_QTY ≤ HOLD_REM the ship comes "
             "entirely from the held reserve; when ALLOC_QTY > HOLD_REM the "
             "difference comes from FNL_Q_REM if available, otherwise the gap is "
             "unfulfilled. In all three cases the tracker treats the hold as "
             "released (HOLD_REM −= ALLOC_QTY, clamped at 0). The conditional "
             "matters when reconciling against actual warehouse picks or "
             "explaining a fulfilled-vs-promised mismatch.")

    add_h(doc, "Stage 8 — Closing the gap (now implemented)", level=2)
    add_para(doc,
             "Previously, if a NEW store F appeared with a TBL need on day 3 "
             "while the 30 held units were still physically at DH24 with no "
             "MASTER_ALC_PEND row, those 30 units counted toward FNL_Q in the "
             "next MSA run, allowing them to be promised to store F as well. "
             "That risk is now structurally closed.")
    tip_box(doc,
            "MSA Step 6.5 (new) reads ARS_NL_TBL_HOLD_TRACKING for IS_CLOSED=0 "
            "rows, maps WERKS → RDC via Master_ALC_INPUT_ST_MASTER, sums "
            "HOLD_REM at (RDC, ARTICLE_NUMBER), and merges into msa_pivot. "
            "Step 7 then computes FNL_Q = max(STK − PEND − HOLD, 0). The 30 "
            "held units no longer enter the next pool, so store F's TBL "
            "request sees only the genuinely-free 90 units (less anything "
            "else pending).")
    add_para(doc,
             "Implementation details and the exact diff are in Section 6.")

    # ── Section 3: Flowchart ───────────────────────────────────────────────
    doc.add_page_break()
    add_h(doc, "3. Flowchart", level=1)
    add_para(doc,
             "Process flow from inward GRN to store delivery. The HOLD branch "
             "(right side) is what's specific to TBL options.", italic=True,
             size=10)
    doc.add_paragraph()

    flow_node(doc, "STAGE 1 · Inward GRN at RDC",
              "ET_STORE_STOCK ← SAP")
    flow_arrow(doc)
    flow_node(doc, "STAGE 2 · MSA Stock Calculation",
              "STK_QTY − PEND_QTY = FNL_Q  →  ARS_MSA_VAR_ART")
    flow_arrow(doc)
    flow_node(doc, "STAGE 3 · Listing Generation",
              "Part 3.6 tags OPT_TYPE; Part 4c computes OPT_MBQ / OPT_MBQ_WH")
    flow_arrow(doc)
    flow_node(doc, "STAGE 4 · Allocator (rule_engine.py)",
              "pool_taken split into SHIP_QTY + HOLD_QTY (TBL only)",
              fill="FFE699", border="BF8F00")
    flow_arrow(doc, "▼ splits ▼")

    flow_branch(doc,
                "ALLOC_QTY", "(= SHIP_QTY) → ARS_ALLOC_WORKING",
                "HOLD_QTY", "(TBL only) → ARS_ALLOC_WORKING")
    flow_arrow(doc)
    flow_branch(doc,
                "STAGE 5 · BDC Export", "Dispatch file → SAP outbound",
                "STAGE 5b · Tracker Insert", "ARS_NL_TBL_HOLD_TRACKING (Step B)")
    flow_arrow(doc)
    flow_branch(doc,
                "STAGE 6 · Warehouse Pick", "MASTER_ALC_PEND populated",
                "STAGE 6b · Held at RDC", "Physically remains in stock")
    flow_arrow(doc)
    flow_branch(doc,
                "STAGE 7 · Store Delivery", "Store stock += SHIP_QTY",
                "STAGE 7b · Next-cycle decrement", "Step A: HOLD_REM −= ALLOC_QTY")
    flow_arrow(doc)
    flow_node(doc,
              "✓ GAP CLOSED — MSA Step 6.5 deducts open HOLD_REM from STK_QTY",
              "FNL_Q = max(STK − PEND − HOLD, 0). Held units no longer offerable to other stores.",
              fill="E2EFDA", border="548235")

    # ── Section 4: state cheat-sheet ──────────────────────────────────────
    doc.add_page_break()
    add_h(doc, "4. What each table holds at each stage", level=1)
    add_para(doc,
             "FNL_Q values reflect the post-Step-6.5 calculation: "
             "FNL_Q = max(STK − PEND − HOLD, 0).", italic=True, size=10)
    make_table(doc,
               ["Stage", "ET_STORE_STOCK (RDC)", "MASTER_ALC_PEND",
                "ARS_MSA_VAR_ART (STK / PEND / HOLD / FNL_Q)",
                "ARS_ALLOC_WORKING", "ARS_NL_TBL_HOLD_TRACKING"],
               [
                   ["1. Inward arrival", "200", "0", "n/a", "n/a", "unchanged"],
                   ["2. After MSA day 1",      "200", "0", "200 / 0 / 0 / 200",
                    "n/a", "unchanged"],
                   ["3-4. After Listing day 1", "200", "0", "(unchanged in this MSA)",
                    "5 rows: ALLOC=16, HOLD=6", "5 new rows: HOLD_REM=6"],
                   ["5. After tracker insert", "200", "0", "—", "unchanged",
                    "5 rows total (30 reserved)"],
                   ["6. After BDC dispatch",   "200 → 120 (when picked)", "+80",
                    "(next MSA)", "unchanged", "unchanged"],
                   ["7. After delivery",       "120", "back to 0",
                    "next MSA: 120 / 0 / 30 / 90", "unchanged", "unchanged"],
                   ["7. After cycle-2 alloc",  "120 → less", "refreshed",
                    "(next MSA, post-deduct)", "ALLOC=4, HOLD=0",
                    "5 rows: HOLD_REM=2 (10 reserved)"],
                   ["7. After cycle-3 alloc",  "less", "refreshed",
                    "(next MSA, post-deduct)", "ALLOC=4, HOLD=0",
                    "5 rows: HOLD_REM=0, IS_CLOSED=1"],
               ],
               col_widths_in=[1.3, 1.1, 1.0, 1.6, 1.2, 1.4])

    # ── Section 5: confusion points ───────────────────────────────────────
    add_h(doc, "5. Four common confusions, answered", level=1)

    add_h(doc, "Q1. ALLOC_QTY > HOLD_QTY — which one ships?", level=3)
    add_para(doc,
             "ALLOC_QTY ships (it equals SHIP_QTY in the allocator). HOLD_QTY "
             "stays at the RDC. Both are stored separately on the same row of "
             "ARS_ALLOC_WORKING. There is no 'winner' — they are different "
             "buckets.")

    add_h(doc, "Q2. What if HOLD_QTY = 0 or HOLD_QTY < REQ?", level=3)
    add_para(doc,
             "For RL and TBC, HOLD_QTY is always 0 by design. For TBL with "
             "hold_days = 0, also 0 (because OPT_MBQ_WH = OPT_MBQ). The "
             "allocator never reads any prior HOLD value to decide what to "
             "allocate; demand is served from the pool, ship gets priority, "
             "hold gets the leftover.")

    add_h(doc, "Q3. Is FNL_Q reduced by HOLD_QTY before the next run?", level=3)
    add_para(doc,
             "Yes — as of this change, MSA Step 6.5 reads "
             "ARS_NL_TBL_HOLD_TRACKING (IS_CLOSED=0) and subtracts SUM(HOLD_REM) "
             "from STK_QTY at the (RDC, ARTICLE_NUMBER) grain, in parallel "
             "to the existing MASTER_ALC_PEND deduction. So FNL_Q now equals "
             "max(STK − PEND − HOLD, 0). See Section 6 for the diff.")

    add_h(doc, "Q4. What does the tracker do now?", level=3)
    add_para(doc,
             "Two jobs: (1) reservation source — MSA Step 6.5 reads it to "
             "deduct held units from the next pool; (2) audit log — Step A "
             "decrements HOLD_REM whenever the same SKU has an ALLOC_QTY this "
             "run (HOLD_REM = MAX(0, HOLD_REM − ALLOC_QTY); the conditional "
             "logic about whether the difference physically came from the "
             "hold or from FNL_Q_REM is conceptual, not encoded in the "
             "decrement SQL). Step A.5 refreshes when the planner re-evaluates "
             "a held SKU, Step B inserts a fresh row for new TBL/NL "
             "allocations.")

    # ── Section 6: gap fix ────────────────────────────────────────────────
    add_h(doc, "6. Closing the gap — what was changed", level=1)

    add_callout(doc, "STATUS →",
                "Option A is now IMPLEMENTED in app/services/msa_service.py "
                "(MSA Step 6.5). Open HOLD_REM is deducted from STK_QTY at "
                "the same grain as PEND_QTY before computing FNL_Q. The "
                "tracker is now a real reservation, not just an audit log.",
                "E2EFDA")

    add_h(doc, "Option A (implemented) — deduct open holds in MSA Step 6.5", level=3)
    add_para(doc,
             "A new helper MSAService._load_open_holds() reads "
             "ARS_NL_TBL_HOLD_TRACKING (IS_CLOSED=0, HOLD_REM>0), maps WERKS → "
             "RDC via Master_ALC_INPUT_ST_MASTER (RDC column probed across "
             "RDC / WAREHOUSE / HUB / WH_CD), and aggregates SUM(HOLD_REM) by "
             "(RDC, ARTICLE_NUMBER). Step 6.5 in calculate() merges this onto "
             "msa_pivot. Step 7 then subtracts both PEND and HOLD:")
    code_block(doc,
               "# msa_service.py — Step 6.5 (new)\n"
               "holds_pivot = self._load_open_holds()\n"
               "if not holds_pivot.empty and \"ARTICLE_NUMBER\" in msa_pivot.columns:\n"
               "    msa_pivot = msa_pivot.merge(\n"
               "        holds_pivot,\n"
               "        left_on=[\"ST_CD\", \"ARTICLE_NUMBER\"],\n"
               "        right_on=[\"RDC\", \"ARTICLE_NUMBER\"],\n"
               "        how=\"left\",\n"
               "    )\n"
               "    msa_pivot[\"HOLD_QTY\"] = msa_pivot[\"HOLD_QTY\"].fillna(0)\n"
               "    msa_pivot.drop(columns=[\"RDC\"], inplace=True, errors=\"ignore\")\n"
               "else:\n"
               "    msa_pivot[\"HOLD_QTY\"] = 0\n\n"
               "# msa_service.py — Step 7 (updated)\n"
               "msa_pivot[\"FNL_Q\"] = np.maximum(\n"
               "    msa_pivot[\"STK_QTY\"] - msa_pivot[\"PEND_QTY\"] - msa_pivot[\"HOLD_QTY\"], 0\n"
               ")",
               lang_label="msa_service.py (already applied)")
    add_para(doc,
             "Effects:")
    add_bullets(doc, [
        "FNL_Q in ARS_MSA_TOTAL / ARS_MSA_GEN_ART / ARS_MSA_VAR_ART now reflects warehouse stock minus held units.",
        "HOLD_QTY appears as a new column on every MSA output table (auto-picked up by msa_result_storage.py since the schema is inferred from the DataFrame).",
        "Allocator (rule_engine.py) needs no change — it reads the already-deducted FNL_Q.",
        "Failure-tolerant: if ARS_NL_TBL_HOLD_TRACKING or Master_ALC_INPUT_ST_MASTER is missing, MSA falls back to HOLD_QTY=0 with a warning log — old behavior is preserved.",
    ])
    warn_box(doc,
             "Sanity check after first run: SUM(HOLD_QTY) in ARS_MSA_TOTAL "
             "should equal SUM(HOLD_REM) in ARS_NL_TBL_HOLD_TRACKING WHERE "
             "IS_CLOSED=0. If they diverge, check the WERKS → RDC mapping in "
             "Master_ALC_INPUT_ST_MASTER for missing or stale rows.")

    add_h(doc, "Option B (alternative) — deduct in rule_engine.py", level=3)
    add_para(doc,
             "Kept here for reference. Smaller change but only visible to "
             "the allocator; not needed now that Option A is in.")
    code_block(doc,
               "-- existing\n"
               "TRY_CAST(V.[FNL_Q] AS FLOAT) AS FNL_Q,\n"
               "TRY_CAST(V.[FNL_Q] AS FLOAT) AS FNL_Q_REM,\n\n"
               "-- patched\n"
               "TRY_CAST(V.[FNL_Q] AS FLOAT) - ISNULL(H.held, 0) AS FNL_Q,\n"
               "TRY_CAST(V.[FNL_Q] AS FLOAT) - ISNULL(H.held, 0) AS FNL_Q_REM,\n"
               "...\n"
               "LEFT JOIN (\n"
               "    SELECT VAR_ART, SZ, SUM(HOLD_REM) AS held\n"
               "    FROM ARS_NL_TBL_HOLD_TRACKING\n"
               "    WHERE IS_CLOSED = 0\n"
               "    GROUP BY VAR_ART, SZ\n"
               ") H ON H.VAR_ART = V.VAR_ART AND H.SZ = V.SZ",
               lang_label="rule_engine.py:357 — alternative patch (NOT applied)")

    # ── Footer ────────────────────────────────────────────────────────────
    add_hr(doc)
    add_para(doc,
             "Source files referenced: "
             "app/api/v1/endpoints/listing.py:2041-2175 (Part 8.6 hold tracker), "
             "app/services/rule_engine.py:1015-1303 (allocator split), "
             "app/services/msa_service.py:461-521 (Step 6 / Step 7).",
             italic=True, size=9, color=(0x80, 0x80, 0x80))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"Wrote {out_path}  ({out_path.stat().st_size:,} bytes)")


if __name__ == "__main__":
    import sys
    out = Path(__file__).parent / "HOLD_QTY_Lifecycle.docx"
    try:
        build_doc(out)
    except PermissionError:
        # File is open in Word — fall back to a versioned name so the run
        # still produces an artifact.
        fallback = out.with_name("HOLD_QTY_Lifecycle_v2.docx")
        build_doc(fallback)
        print(f"NOTE: original was locked (open in Word). Wrote {fallback} instead.",
              file=sys.stderr)
